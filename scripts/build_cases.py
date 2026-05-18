from __future__ import annotations

import csv
import re
import subprocess
import unicodedata
from pathlib import Path

from docx import Document
from pypdf import PdfReader


BENCH_DIR = Path(__file__).resolve().parents[1]
WORKSPACE_DIR = BENCH_DIR.parent
RAW_RESUME_DIR = WORKSPACE_DIR / "测试简历"
DATA_DIR = BENCH_DIR / "data"
ANON_RESUME_DIR = DATA_DIR / "resumes_anon"
JD_DIR = DATA_DIR / "jds"


REAL_NAMES = [
    "贺常志",
    "He Changzhi",
    "刘楚焕",
    "Lucky",
    "丁屹涵",
    "巩山虹",
    "樊瑞琪",
    "洪永昊",
    "洪赫",
    "王旭",
    "修丹阳",
    "陈召基",
    "陈宇锐",
    "黄天立",
    "黎前",
    "大春",
]


JD_LIBRARY = {
    "senior_java_backend": {
        "source": "https://remote-china.com/jobs/1056",
        "text": """岗位：Java 后端开发工程师（营销平台/中台方向）

岗位职责：
1. 负责营销活动、营销中台相关业务的后端模块开发与迭代，将业务规则转化为稳定的技术方案。
2. 参与订单、活动配置、用户指标统计、数据采集与清洗等模块建设。
3. 参与系统性能分析、稳定性治理、代码评审和技术方案评审。
4. 与产品、前端、测试、数据团队协作，保障需求按期交付。

任职要求：
1. 6 年以上 Java 开发经验，熟悉 Spring Boot、Spring Cloud、MyBatis 等主流框架。
2. 熟悉 MySQL、Redis、Kafka 等常用数据库和中间件，理解常见性能优化方法。
3. 有营销活动、增长平台、交易系统或企业级业务系统经验优先。
4. 具备良好的代码质量意识、问题排查能力和团队协作能力。""",
    },
    "ai_product_manager": {
        "source": "https://jobs.lenovo.com/en_US/careers/JobDetail/AI/70467",
        "text": """岗位：AI 产品经理（大模型应用方向）

岗位职责：
1. 负责 AI 产品需求调研、场景分析、原型设计和 PRD 撰写，推动产品从需求到上线。
2. 与算法、研发、设计、运营团队协同，推进 AI 能力在业务流程中的落地。
3. 跟踪大模型、智能体、知识库等方向的产品形态，持续优化用户体验和业务指标。
4. 基于用户反馈和数据分析，迭代产品功能与运营策略。

任职要求：
1. 具备产品经理基础能力，能独立完成需求拆解、流程设计和跨团队推进。
2. 熟悉 AI 产品基本概念，有大模型、RAG、Agent、模型评测或 Prompt 设计经验者优先。
3. 具备优秀的信息整理、逻辑表达、项目推进和用户洞察能力。
4. 对 AI 产品长期保持关注，能快速学习新工具和新场景。""",
    },
    "education_delivery": {
        "source": "https://www.wendangwuyou.com/fw/133665.html",
        "text": """岗位：课程交付/学习运营老师

岗位职责：
1. 负责线上课程交付全流程，包括开课准备、学员沟通、直播支持、学习提醒和结营复盘。
2. 跟踪学习数据，推动完课率、作业提交率和学员满意度提升。
3. 协同教研、讲师、销售和社群团队，优化课程流程与服务体验。
4. 沉淀课程交付 SOP、常见问题和学员反馈，支持后续课程迭代。

任职要求：
1. 具备教育培训、用户运营、课程运营或客户服务经验。
2. 沟通表达清晰，有较强执行力、责任心和服务意识。
3. 能使用表格工具进行基础数据统计和复盘。
4. 有直播课、训练营、陪跑营或成人教育项目经验优先。""",
    },
    "battery_materials": {
        "source": "https://bebee.com/cn/jobs/atl-nd--ss-cn-9etj0",
        "text": """岗位：新能源材料研发工程师（电化学/电池材料方向）

岗位职责：
1. 参与电池材料、电化学体系或可持续材料方向的实验设计、样品制备与性能测试。
2. 整理实验数据，分析材料结构、性能和工艺参数之间的关系。
3. 跟踪新能源材料、电化学储能、绿色能源等方向文献和技术进展。
4. 撰写实验记录、技术报告和阶段性总结，协助跨团队沟通。

任职要求：
1. 材料、化工、电化学、能源或相关专业硕士及以上学历。
2. 熟悉材料表征、电化学基础知识和实验室安全规范。
3. 具备英文文献阅读、数据整理和实验复盘能力。
4. 有电池材料、催化、储能或可持续发展相关项目经验优先。""",
    },
    "securities_institutional": {
        "source": "https://www.zhaopin.com/jobdetail/CC600054120J40686481806.htm",
        "text": """岗位：证券机构业务经理

岗位职责：
1. 开发与维护私募基金、银行、信托、保险、券商等金融机构客户。
2. 协助推进机构客户服务、产品路演、业务培训和重点客户关系维护。
3. 跟踪资本市场动态和客户需求，输出业务机会和客户经营建议。
4. 遵守证券行业合规要求，完成分支机构下达的业务目标。

任职要求：
1. 金融、经济、管理等相关专业背景，熟悉证券市场和金融产品。
2. 具备机构业务、营业部销售管理、金融客户服务或渠道拓展经验。
3. 沟通表达、抗压能力、客户服务意识和合规意识强。
4. 通过证券行业专业人员水平评价测试或基金从业资格者优先。""",
    },
    "data_analyst": {
        "source": "https://www.glassdoor.com/job-listing/senior-data-analyst-sql-and-tableau-ttw-solutions-JV_IC3728413_KO0%2C35_KE36%2C49.htm?jl=1009822224073",
        "text": """岗位：数据分析师（SQL/Python/BI）

岗位职责：
1. 负责业务数据提取、清洗、校验、分析和可视化，支持运营与产品决策。
2. 建立常规数据报表和分析看板，跟踪核心指标变化并解释异常。
3. 使用 SQL、Python 或表格工具完成数据处理、自动化分析和结论沉淀。
4. 与业务、产品、研发团队沟通分析口径，推动数据分析结果落地。

任职要求：
1. 统计、数学、计算机、商业分析、经济金融等相关专业背景。
2. 熟悉 SQL，掌握 Python/R/Excel 等至少一种分析工具。
3. 具备清晰的数据逻辑、指标意识和报告表达能力。
4. 有 Tableau、Power BI、Looker 或数据仓库经验优先。""",
    },
    "ai_content_ops": {
        "source": "https://career.nankai.edu.cn/correcruit/content/id/113894.html",
        "text": """岗位：AI 内容运营/产品运营

岗位职责：
1. 结合产品特性与用户洞察，策划图文、视频或社区内容，提升点击、互动和转化。
2. 跟踪阅读量、分享率、留存、转化路径等内容数据，持续优化内容策略。
3. 参与 AI 工具或内容产品的需求反馈、用户调研和运营活动执行。
4. 沉淀内容选题、发布节奏、复盘报告和运营方法。

任职要求：
1. 具备内容运营、产品运营、用户研究、媒体传播或数据分析相关经历。
2. 对 AI 工具、内容平台和创作者生态有兴趣，学习能力强。
3. 能进行基础数据分析，表达清晰，有良好的协作和执行能力。
4. 有校园媒体、社群运营、产品实习或内容增长经验优先。""",
    },
    "cv_algorithm": {
        "source": "https://www.saikr.com/practice/job/11711",
        "text": """岗位：计算机视觉算法工程师

岗位职责：
1. 参与目标检测、图像分割、跟踪、关键点等计算机视觉算法研发。
2. 负责数据处理、模型训练、测试评估和算法效果分析，支持业务场景落地。
3. 跟踪视觉算法前沿进展，整理实验结果、技术文档和复盘报告。
4. 与工程、产品或硬件团队协作，推进算法部署和迭代。

任职要求：
1. 计算机、人工智能、自动化、数据科学等相关专业背景。
2. 熟悉机器学习、深度学习和计算机视觉基础，掌握 Python。
3. 有目标检测、分割、跟踪、图像分类或机器人感知项目经验。
4. 熟悉 PyTorch/TensorFlow/OpenCV、模型部署或竞赛论文经验优先。""",
    },
    "junior_backend": {
        "source": "https://www.ziprecruiter.com/c/ITR-Group/Job/Backend-Engineer-%28Java-Spring-Kafka%29/-in-Minneapolis%2CMN?jid=bdc7e836b47390be",
        "text": """岗位：后端开发工程师（Java/实习或校招）

岗位职责：
1. 参与业务系统后端模块开发、接口联调、单元测试和技术文档编写。
2. 根据需求完成模块设计、数据表设计、错误处理和基础性能优化。
3. 使用 Git 进行版本协作，配合测试、前端和产品完成迭代交付。
4. 参与代码评审、问题排查和项目复盘。

任职要求：
1. 计算机相关专业，熟悉 Java 基础、面向对象编程和常见数据结构。
2. 了解 Spring Boot、SQL、Git、Linux 或 CI/CD 基础。
3. 有课程项目、实习项目或开源协作经验。
4. 具备良好的学习能力、工程规范意识和团队沟通能力。""",
    },
    "fpga_engineer": {
        "source": "https://bebee.com/cn/jobs/fpga--ss-cn-ln2x4z",
        "text": """岗位：FPGA 工程师

岗位职责：
1. 参与 FPGA 方案设计、模块开发、仿真验证、板级调试和文档编写。
2. 根据产品需求完成接口、时序、数据通路或信号处理相关逻辑设计。
3. 配合硬件、软件和测试团队定位问题，推动设计迭代。
4. 维护设计文档、测试记录和版本交付资料。

任职要求：
1. 电子信息、通信工程、自动化、集成电路等相关专业背景。
2. 熟悉数字电路、通信原理、信号处理和 FPGA 基础开发流程。
3. 了解 Verilog/VHDL、Vivado/Quartus、ModelSim 等工具者优先。
4. 具备清晰逻辑思维、学习能力和团队合作意识。""",
    },
    "admin_specialist": {
        "source": "https://bebee.com/cn/jobs/ss-cn-1spij8j",
        "text": """岗位：行政专员/综合办公室专员

岗位职责：
1. 负责会议组织、会务支持、公文流转、档案整理和日常行政事务。
2. 管理办公用品、固定资产、采购台账和费用报销材料。
3. 协助部门完成制度执行、跨部门沟通、来访接待和内部通知。
4. 整理行政数据和工作记录，保障办公室流程规范高效。

任职要求：
1. 具备行政、文秘、综合办公室、政府机关或企业职能支持经验。
2. 公文写作、档案管理、会议管理和办公软件使用能力良好。
3. 做事严谨细致，责任心强，沟通协调和保密意识好。
4. 有招聘协助、流程管理或固定资产管理经验优先。""",
    },
    "finance_assistant": {
        "source": "https://www.gdnybank.com/upload/ew/20231229141926282.pdf",
        "text": """岗位：金融业务助理/理财顾问助理

岗位职责：
1. 协助客户资料整理、产品材料准备、客户沟通记录和业务流程跟进。
2. 支持金融产品、市场信息和客户需求的整理分析。
3. 配合团队完成客户服务、活动执行、基础报表和合规材料。
4. 学习银行、证券、基金、保险等金融产品知识，提升客户服务能力。

任职要求：
1. 金融、投资、经济、财务管理等相关专业背景。
2. 熟悉基础金融市场知识，具备良好的数字敏感度和服务意识。
3. 沟通表达清晰，工作细致，能按规范处理客户资料。
4. 具备证券、基金、银行从业相关证书或实习经验优先。""",
    },
    "ui_visual_designer": {
        "source": "https://www.hackquest.io/jobs/338990",
        "text": """岗位：UI/UX 视觉设计师

岗位职责：
1. 负责产品界面、运营活动、品牌视觉和推广物料的设计输出。
2. 与产品、开发、运营团队协作，完成原型、视觉稿、组件规范和交付标注。
3. 基于业务目标和用户体验优化页面结构、视觉层级和交互细节。
4. 关注设计趋势和 AI 设计工具，提升视觉效率和品牌一致性。

任职要求：
1. 具备 UI、视觉、品牌或平面设计经验，审美和版式能力扎实。
2. 熟悉 Figma、Sketch、PS、AI 等设计工具，有组件化设计意识。
3. 能清晰表达设计思路，具备跨团队沟通和按期交付能力。
4. 有完整产品项目、动效、图标、品牌延展或 AIGC 设计经验优先。""",
    },
    "ai_backend_intern": {
        "source": "https://pic.bankofchina.com/bocappd/appform/202507/P020250702344144547288.pdf",
        "text": """岗位：AI 后端开发实习生

岗位职责：
1. 参与 AI 应用后端模块开发，包括数据处理、接口开发、任务调度和服务联调。
2. 协助完成模型调用、Prompt 配置、日志记录、测试验证和问题排查。
3. 与算法、产品和前端同学协作，推动 AI 功能原型验证和迭代。
4. 编写技术文档、测试用例和项目复盘材料。

任职要求：
1. 计算机、软件工程、人工智能等相关专业在读。
2. 熟悉 Java 或 Python，了解数据库、Git、Linux 和基础后端开发流程。
3. 对大模型应用、RAG、Agent、Prompt 工程或向量数据库有兴趣。
4. 有 AI 项目、后端项目或算法实践经验优先。""",
    },
    "chemical_process": {
        "source": "https://startup.jobs/e-chem-design-engineer-cell-electrode-rivian-automotive-2-2650064",
        "text": """岗位：化工/电化学工艺工程师

岗位职责：
1. 参与电化学体系、材料工艺或化工流程相关实验设计与数据分析。
2. 协助推进实验方案、参数优化、样品测试和技术报告撰写。
3. 结合文献和实验结果，提出材料性能、工艺稳定性或可持续性改进建议。
4. 与研发、测试、供应链或质量团队沟通，支持项目阶段性推进。

任职要求：
1. 化学工程、材料、能源、电化学或相关专业背景。
2. 熟悉化工热力学、传递过程、反应工程或电化学基础。
3. 具备英文文献阅读、实验记录、数据处理和安全规范意识。
4. 有新能源、储能、可持续材料、工艺放大或流程模拟经验优先。""",
    },
    "course_ops_manager": {
        "source": "https://career.cuhk.edu.cn/attachment/careercuhk/ueditor/file/20250612/7204_%E4%BA%A7%E5%93%81%E8%BF%90%E8%90%A5-%E5%AE%9E%E4%B9%A0%E5%B2%97%E4%BD%8D.pdf",
        "text": """岗位：课程运营/训练营运营负责人

岗位职责：
1. 负责训练营或线上课程的排期、招生协同、开课执行、社群督学和结营复盘。
2. 搭建课程运营流程，监控报名、到课、作业、完课、满意度等关键指标。
3. 协同讲师、教研、销售、客服和社群团队，提升学员体验与课程交付质量。
4. 基于数据和反馈迭代课程内容、运营动作和服务 SOP。

任职要求：
1. 具备教育行业、知识付费、训练营、用户运营或项目运营经验。
2. 有较强项目管理、跨部门协作、复盘分析和现场执行能力。
3. 能使用表格或数据工具跟踪指标，输出运营复盘。
4. 有千人级社群、直播课或成人学习项目经验优先。""",
    },
}


CASES = [
    {
        "id": "case_04",
        "description": "Case 04 - 9年 Java 后端 -> 营销中台高级后端",
        "candidate_name": "候选人04",
        "source_file": "个人简历_贺常志.pdf",
        "jd": "senior_java_backend",
        "category": "backend_java",
        "forbidden_terms": "高并发,秒杀,Flink,Hadoop,Spark,ES,MongoDB,自研营销中台,营销中台架构,架构设计,主导,核心模块,监控体系",
    },
    {
        "id": "case_05",
        "description": "Case 05 - 新闻学非应届 -> AI 产品经理",
        "candidate_name": "候选人05",
        "source_file": "于大春.pdf",
        "jd": "ai_product_manager",
        "category": "ai_product_transition",
        "forbidden_terms": "RAG,AI Agent,模型微调,LangChain,Dify,Coze,模型评测,Prompt工程,从0到1大模型项目,知识库",
    },
    {
        "id": "case_06",
        "description": "Case 06 - 7年教培运营 -> 课程交付老师",
        "candidate_name": "候选人06",
        "source_file": "刘楚焕+应聘交付老师+7年教培行业经验+13902267823.pdf",
        "jd": "education_delivery",
        "category": "education_delivery",
        "forbidden_terms": "NPS,企业培训,SaaS,客户成功,CRM,私域增长,留存率,客户分层,满意度调研",
    },
    {
        "id": "case_07",
        "description": "Case 07 - 化工硕士 -> 新能源材料研发",
        "candidate_name": "候选人07",
        "source_file": "宾夕法尼亚大学_丁屹涵.pdf",
        "jd": "battery_materials",
        "category": "materials_energy",
        "forbidden_terms": "正极材料,锂电,电芯,扣式电池,SEM,XRD,DOE,Minitab,JMP,量产导入",
    },
    {
        "id": "case_08",
        "description": "Case 08 - 证券营业部助理 -> 机构业务经理",
        "candidate_name": "候选人08",
        "source_file": "巩山虹 个人中文简历(1).docx",
        "jd": "securities_institutional",
        "category": "finance_securities",
        "forbidden_terms": "私募基金,上市公司资源,投行业务,资产托管,PB业务,量化策略,合规风控体系,两融业务",
    },
    {
        "id": "case_09",
        "description": "Case 09 - 分析学硕士 -> 数据分析师",
        "candidate_name": "候选人09",
        "source_file": "樊瑞琪简历.pdf",
        "jd": "data_analyst",
        "category": "data_analysis",
        "forbidden_terms": "Tableau,Power BI,Snowflake,数据仓库,ETL,数据建模,Looker,BI报表",
    },
    {
        "id": "case_10",
        "description": "Case 10 - 影视文学本科 -> AI 内容运营",
        "candidate_name": "候选人10",
        "source_file": "洪永昊-中文简历.pdf",
        "jd": "ai_content_ops",
        "category": "content_ops",
        "forbidden_terms": "AIGC,提示词,内容SOP,转化率,留存率,创作者生态,SQL,商业化",
    },
    {
        "id": "case_11",
        "description": "Case 11 - 机器人感知实习 -> 计算机视觉算法",
        "candidate_name": "候选人11",
        "source_file": "洪赫—数据、AI算法工程师.pdf",
        "jd": "cv_algorithm",
        "category": "cv_algorithm",
        "forbidden_terms": "CVPR,ICCV,Transformer,BEV,Diffusion,Docker,TensorRT,多模态,大模型",
    },
    {
        "id": "case_12",
        "description": "Case 12 - CS 本科实习 -> Java 后端校招",
        "candidate_name": "候选人12",
        "source_file": "王旭简历(1).pdf",
        "jd": "junior_backend",
        "category": "junior_backend",
        "forbidden_terms": "Spring Cloud,Dubbo,Kafka,Redis,高并发,微服务架构,生产环境,性能优化,分布式事务",
    },
    {
        "id": "case_13",
        "description": "Case 13 - 电子信息硕士 -> FPGA 工程师",
        "candidate_name": "候选人13",
        "source_file": "简历FPGA(1).pdf",
        "jd": "fpga_engineer",
        "category": "fpga_hardware",
        "forbidden_terms": "VHDL,Vivado,SystemVerilog,UVM,DDR,PCIE,SerDes,AXI,时序约束,高速接口",
    },
    {
        "id": "case_14",
        "description": "Case 14 - 政府办公室职员 -> 行政专员",
        "candidate_name": "候选人14",
        "source_file": "陈召基简历(1).pdf",
        "jd": "admin_specialist",
        "category": "admin",
        "forbidden_terms": "招聘全流程,薪酬绩效,固定资产系统,ERP,预算管理,董事会秘书,政府关系,OA系统搭建",
    },
    {
        "id": "case_15",
        "description": "Case 15 - 投资理财应届 -> 金融业务助理",
        "candidate_name": "候选人15",
        "source_file": "陈宇锐的简历.pdf",
        "jd": "finance_assistant",
        "category": "finance_assistant",
        "forbidden_terms": "基金从业资格,证券从业资格,理财规划,资产配置,客户开发,销售业绩,私募基金,投顾服务",
    },
    {
        "id": "case_16",
        "description": "Case 16 - 视觉/UI 设计师 -> UI/UX 视觉设计",
        "candidate_name": "候选人16",
        "source_file": "黄天立-简历（2024）.pdf",
        "jd": "ui_visual_designer",
        "category": "ui_design",
        "forbidden_terms": "Figma,Sketch,组件库,设计系统,0到1产品,Web3,动效设计,用户研究,交互原型",
    },
    {
        "id": "case_17",
        "description": "Case 17 - English Java resume -> backend engineer",
        "candidate_name": "Candidate17",
        "source_file": "He Changzhi.pdf",
        "jd": "senior_java_backend",
        "category": "backend_java_en",
        "forbidden_terms": "microservices architecture,high concurrency,system monitoring,Elasticsearch,MongoDB,Flink,Hadoop,Spark,technical leadership,architect",
    },
    {
        "id": "case_18",
        "description": "Case 18 - CS 本科项目 -> AI 后端实习",
        "candidate_name": "候选人18",
        "source_file": "王旭.doc",
        "jd": "ai_backend_intern",
        "category": "ai_backend_intern",
        "forbidden_terms": "RAG,LangChain,向量数据库,大模型微调,Agent,Prompt工程,模型评测,线上AI服务,CUDA",
    },
    {
        "id": "case_19",
        "description": "Case 19 - 化工硕士 -> 化工/电化学工艺",
        "candidate_name": "候选人19",
        "source_file": "宾夕法尼亚大学_丁屹涵.pdf",
        "jd": "chemical_process",
        "category": "chemical_process",
        "forbidden_terms": "P&ID,中试放大,工艺包,反应器设计,安全评价,量产,COMSOL,DOE,工艺验证",
    },
    {
        "id": "case_20",
        "description": "Case 20 - 教培运营 -> 训练营运营负责人",
        "candidate_name": "候选人20",
        "source_file": "刘楚焕+应聘交付老师+7年教培行业经验+13902267823.pdf",
        "jd": "course_ops_manager",
        "category": "course_ops",
        "forbidden_terms": "SaaS客户成功,CRM,企业客户,交付SOP,客诉处理,NPS,私域转化,销售漏斗,学习中台,企业培训",
    },
]


EXISTING_CASES = [
    {
        "id": "case_01",
        "description": "Case 01 - 初级前端（2-3年经验）→ 高级前端（AI 创业公司，要求 4 年+）",
        "candidate_name": "张明远",
        "resume": "file://data/case_01_resume.txt",
        "jd": "file://data/case_01_jd.txt",
        "category": "frontend_ai",
        "forbidden_terms": "TypeScript,Next.js,Nuxt,SSR,Webpack,Vite,AI 产品,大模型,开源",
        "source": "synthetic",
    },
    {
        "id": "case_02",
        "description": "Case 02 - 5 年电商 PM → AI 产品经理（要求 LLM 经验 + prompt engineering）",
        "candidate_name": "李婉清",
        "resume": "file://data/case_02_resume.txt",
        "jd": "file://data/case_02_jd.txt",
        "category": "pm_ai_transition",
        "forbidden_terms": "LLM,大语言模型,prompt engineering,Prompt 设计,模型评测,benchmark,RAG,Agent,多模态,AI 产品,智能客服,AI 写作,技术博客",
        "source": "synthetic",
    },
    {
        "id": "case_03",
        "description": "Case 03 - 应届统计学本科 → 数据分析师（要求 SQL 中阶 + 可视化工具）",
        "candidate_name": "王思宇",
        "resume": "file://data/case_03_resume.txt",
        "jd": "file://data/case_03_jd.txt",
        "category": "data_analysis_entry",
        "forbidden_terms": "Tableau,FineBI,Quicksight,业务分析报告",
        "source": "synthetic",
    },
]


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        doc = Document(str(path))
        parts = [para.text for para in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    if suffix == ".doc":
        proc = subprocess.run(
            ["textutil", "-convert", "txt", "-stdout", str(path)],
            check=False,
            capture_output=True,
            text=True,
            timeout=30,
        )
        return proc.stdout
    raise ValueError(f"Unsupported resume type: {path}")


def normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\u200b", "").replace("\uf0b7", "-")
    text = re.sub(r"[\uf06c\uf09f\uf06e\uf07c\uf0d8]", "-", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def anonymize(text: str, candidate_name: str) -> str:
    text = normalize_text(text)
    for real_name in REAL_NAMES:
        text = text.replace(real_name, candidate_name)

    text = re.sub(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", "[EMAIL]", text)
    text = re.sub(r"(?<!\d)(?:\+?86[-\s]?)?1[3-9]\d[-\s]?\d{4}[-\s]?\d{4}(?!\d)", "[PHONE]", text)
    text = re.sub(r"(?<!\d)\d{3,4}[-\s]\d{3,4}[-\s]\d{3,4}(?!\d)", "[PHONE]", text)
    text = re.sub(r"(?<!\d)\d{10,}(?!\d)", "[PHONE]", text)
    text = re.sub(r"(微信|Wechat|WeChat|QQ|手机号|联系电话|电话|Phone|Tel)[：:\s]*[A-Za-z0-9_+\- ]{5,}", r"\1：[已脱敏]", text, flags=re.I)
    text = re.sub(r"(邮箱|Email|E-mail)[：:\s]*\S+", r"\1：[EMAIL]", text, flags=re.I)
    text = re.sub(r"(地址|户籍地址|现居住地|住址|Address)[：:\s]*[^\n|]+", r"\1：[已脱敏]", text, flags=re.I)
    text = re.sub(r"(出生年月|出生日期|生日|Date of Birth)[：:\s]*[0-9./年月日 -]+", r"\1：[已脱敏]", text, flags=re.I)
    text = re.sub(r"(年龄|Age)[：:\s]*\d{1,2}", r"\1：[已脱敏]", text, flags=re.I)
    text = re.sub(r"\[EMAIL[：:]\[EMAIL\]", "[EMAIL]", text)

    if candidate_name.startswith("Candidate"):
        header = f"Name: {candidate_name}\n"
    else:
        header = f"姓名：{candidate_name}\n"
    if candidate_name not in text[:200]:
        text = header + text
    return text


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.strip() + "\n", encoding="utf-8")


def raw_source_label(case_id: str, source_file: str) -> str:
    return f"{case_id}_raw{Path(source_file).suffix.lower()}"


def build() -> None:
    if not RAW_RESUME_DIR.exists():
        raise FileNotFoundError(f"Raw resume directory not found: {RAW_RESUME_DIR}")

    ANON_RESUME_DIR.mkdir(parents=True, exist_ok=True)
    JD_DIR.mkdir(parents=True, exist_ok=True)

    rows = []
    source_rows = []

    for case in EXISTING_CASES:
        rows.append(
            {
                "__description": case["description"],
                "candidate_name": case["candidate_name"],
                "resume": case["resume"],
                "jd": case["jd"],
                "forbidden_terms": case["forbidden_terms"],
                "__metadata:case_id": case["id"],
                "__metadata:category": case["category"],
                "__metadata:jd_source_url": case["source"],
                "__metadata:raw_resume_source": "synthetic",
                "__metadata:privacy": "synthetic",
            }
        )

    for case in CASES:
        raw_path = RAW_RESUME_DIR / case["source_file"]
        raw_text = extract_text(raw_path)
        anon_text = anonymize(raw_text, case["candidate_name"])
        if len(re.sub(r"\s+", "", anon_text)) < 250:
            raise ValueError(f"Extracted text is too short for {raw_path}")

        resume_ref = f"file://data/resumes_anon/{case['id']}_resume.txt"
        jd_ref = f"file://data/jds/{case['id']}_jd.txt"
        jd = JD_LIBRARY[case["jd"]]

        write_text(ANON_RESUME_DIR / f"{case['id']}_resume.txt", anon_text)
        write_text(JD_DIR / f"{case['id']}_jd.txt", jd["text"])

        rows.append(
            {
                "__description": case["description"],
                "candidate_name": case["candidate_name"],
                "resume": resume_ref,
                "jd": jd_ref,
                "forbidden_terms": case["forbidden_terms"],
                "__metadata:case_id": case["id"],
                "__metadata:category": case["category"],
                "__metadata:jd_source_url": jd["source"],
                "__metadata:raw_resume_source": raw_source_label(case["id"], case["source_file"]),
                "__metadata:privacy": "personal_name_phone_email_address_redacted",
            }
        )
        source_rows.append(
            {
                "case_id": case["id"],
                "raw_resume_source": raw_source_label(case["id"], case["source_file"]),
                "anonymized_resume": resume_ref,
                "jd": jd_ref,
                "jd_source_url": jd["source"],
                "category": case["category"],
            }
        )

    cases_path = DATA_DIR / "cases.csv"
    with cases_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)

    sources_path = DATA_DIR / "case_sources.csv"
    with sources_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(source_rows[0].keys()))
        writer.writeheader()
        writer.writerows(source_rows)

    print(f"Wrote {len(rows)} cases to {cases_path}")
    print(f"Wrote anonymized resumes to {ANON_RESUME_DIR}")
    print(f"Wrote JD files to {JD_DIR}")


if __name__ == "__main__":
    build()
